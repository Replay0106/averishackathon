"""
sdoc_loader.py — Multi-format attachment loader and inbox reader for NavisAI SDOC Hackathon pipeline.

Supports:
  - .txt (UTF-8, Latin-1, CP1252)
  - .pdf (pypdf with corruption trapping and scanned-detection)
  - .docx (python-docx + native zipfile fallback)
  - .xlsx (openpyxl + native zipfile fallback)

Identifies document content types (SI, BL, INVOICE, PACKING_LIST, COO) to detect disguised files.
"""

import json
import os
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from pypdf import PdfReader
    from pypdf.errors import PdfStreamError
except ImportError:
    PdfReader = None
    PdfStreamError = Exception

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import docx
except ImportError:
    docx = None


@dataclass
class AttachmentData:
    path: str
    filename: str
    extension: str
    text: str = ""
    raw_bytes: bytes = b""
    is_corrupted: bool = False
    is_scanned: bool = False
    error_message: Optional[str] = None
    detected_doc_type: str = "UNKNOWN"  # "SI", "BL", "INVOICE", "PACKING_LIST", "COO", "UNKNOWN"


class InboxLoader:
    def __init__(self, bundle_dir: str):
        self.bundle_dir = Path(bundle_dir)
        self.inbox_dir = self.bundle_dir / "inbox"
        self.attachments_dir = self.bundle_dir / "attachments"
        self.injected_emails: Dict[str, Dict[str, Any]] = {}
        self.injected_attachments: Dict[str, AttachmentData] = {}

    def inject_email(self, email_id: str, email_data: Dict[str, Any], attachments: Optional[Dict[str, AttachmentData]] = None):
        """Inject a dynamic or simulated email directly into this loader."""
        self.injected_emails[email_id] = email_data
        if attachments:
            self.injected_attachments.update(attachments)

    def get_email_ids(self) -> List[str]:
        files = sorted(self.inbox_dir.glob("*.json")) if self.inbox_dir.is_dir() else []
        file_ids = [f.stem for f in files]
        # Injected emails appear first sorted in descending order (highest counter/newest first)
        def _sort_key(eid: str):
            digits = re.findall(r"\d+", eid)
            return int(digits[-1]) if digits else 0

        injected = sorted(self.injected_emails.keys(), key=_sort_key, reverse=True)
        return injected + [fid for fid in file_ids if fid not in self.injected_emails]

    def load_emails(self) -> List[Dict[str, Any]]:
        emails = []
        for eid in self.get_email_ids():
            emails.append(self.get_email(eid))
        return emails

    def get_email(self, email_id: str) -> Dict[str, Any]:
        if email_id in self.injected_emails:
            return self.injected_emails[email_id]
        path = self.inbox_dir / f"{email_id}.json"
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return json.load(f)

    def load_attachment(self, rel_path: str) -> AttachmentData:
        filename = Path(rel_path).name
        if filename in self.injected_attachments:
            return self.injected_attachments[filename]
        if rel_path in self.injected_attachments:
            return self.injected_attachments[rel_path]

        full_path = self.bundle_dir / rel_path
        if not full_path.exists():
            for alt_base in [
                Path(tempfile.gettempdir()) / "navis_datasets" / "gmail_live",
                self.bundle_dir.parent / "datasets" / "gmail_live",
            ]:
                alt_path = alt_base / rel_path
                if alt_path.exists():
                    full_path = alt_path
                    break

        ext = full_path.suffix.lower()

        if not full_path.exists():
            return AttachmentData(
                path=rel_path,
                filename=filename,
                extension=ext,
                is_corrupted=True,
                error_message=f"File not found: {rel_path}",
                detected_doc_type="UNKNOWN"
            )

        try:
            raw_bytes = full_path.read_bytes()
        except Exception as e:
            return AttachmentData(
                path=rel_path,
                filename=filename,
                extension=ext,
                is_corrupted=True,
                error_message=f"Read error: {e}",
                detected_doc_type="UNKNOWN"
            )

        att = AttachmentData(
            path=rel_path,
            filename=filename,
            extension=ext,
            raw_bytes=raw_bytes
        )

        if ext == ".txt":
            self._parse_txt(att, raw_bytes)
        elif ext == ".pdf":
            self._parse_pdf(att, full_path, raw_bytes)
        elif ext == ".docx":
            self._parse_docx(att, full_path, raw_bytes)
        elif ext in (".xlsx", ".xls"):
            self._parse_xlsx(att, full_path, raw_bytes)
        else:
            self._parse_txt(att, raw_bytes)

        att.detected_doc_type = self._detect_doc_type(att.text, filename)
        return att

    def _parse_txt(self, att: AttachmentData, raw_bytes: bytes) -> None:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                att.text = raw_bytes.decode(enc)
                return
            except UnicodeDecodeError:
                continue
        att.text = raw_bytes.decode("utf-8", errors="replace")

    def _parse_pdf(self, att: AttachmentData, full_path: Path, raw_bytes: bytes) -> None:
        if PdfReader is None:
            if not raw_bytes.startswith(b"%PDF") or len(raw_bytes) < 1000:
                att.is_corrupted = True
                att.error_message = "Corrupt PDF stream or pypdf not installed"
                return
            if b"/Image" in raw_bytes or b"/XObject" in raw_bytes:
                att.is_scanned = True
                return
            text_chunks = re.findall(rb"\(([^)]+)\)", raw_bytes)
            if text_chunks:
                att.text = "\n".join(c.decode("latin-1", errors="ignore") for c in text_chunks)
                if len(att.text.strip()) < 50:
                    att.is_scanned = True
            else:
                att.is_scanned = True
            return

        try:
            reader = PdfReader(str(full_path))
            pages_text = []
            for p in reader.pages:
                t = p.extract_text() or ""
                pages_text.append(t)
            combined = "\n".join(pages_text).strip()
            att.text = combined
            if len(combined) < 50:
                att.is_scanned = True
        except Exception as e:
            att.is_corrupted = True
            att.error_message = f"PDF parsing error: {e}"

    def _parse_docx(self, att: AttachmentData, full_path: Path, raw_bytes: bytes) -> None:
        if docx is not None:
            try:
                doc = docx.Document(str(full_path))
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                for t in doc.tables:
                    for row in t.rows:
                        row_text = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_text:
                            paragraphs.append(" : ".join(row_text))
                att.text = "\n".join(paragraphs).strip()
                return
            except Exception:
                pass

        # Fallback XML parsing
        try:
            with zipfile.ZipFile(full_path) as z:
                xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            texts = []
            for elem in tree.iter():
                if elem.tag.endswith("}t") and elem.text:
                    texts.append(elem.text)
                elif elem.tag.endswith("}p"):
                    texts.append("\n")
            att.text = "".join(texts).strip()
        except Exception as e:
            att.is_corrupted = True
            att.error_message = f"DOCX parsing error: {e}"

    def _parse_xlsx(self, att: AttachmentData, full_path: Path, raw_bytes: bytes) -> None:
        if openpyxl is not None:
            try:
                wb = openpyxl.load_workbook(str(full_path), data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                        if cells:
                            lines.append(" : ".join(cells))
                att.text = "\n".join(lines).strip()
                return
            except Exception:
                pass

        # Fallback XML parsing
        try:
            with zipfile.ZipFile(full_path) as z:
                texts = []
                for name in z.namelist():
                    if name.startswith("xl/worksheets/sheet") or name == "xl/sharedStrings.xml":
                        tree = ET.fromstring(z.read(name))
                        for elem in tree.iter():
                            if elem.tag.endswith("}t") and elem.text:
                                texts.append(elem.text)
                att.text = "\n".join(texts).strip()
        except Exception as e:
            att.is_corrupted = True
            att.error_message = f"XLSX parsing error: {e}"

    def _detect_doc_type(self, text: str, filename: str) -> str:
        """Determines document type based on content first, then filename fallback. A cover sheet or blank sample in
        front of the document title ("COMMERCIAL INVOICE: NOT INCLUDED") does not decide the type."""
        from sdoc_textnorm import strip_cover

        upper = strip_cover(text.replace("\r\n", "\n").replace("\r", "\n")).upper()
        # Content header checks
        if "COMMERCIAL INVOICE" in upper or "INVOICE NO" in upper:
            return "INVOICE"
        if "PACKING LIST" in upper or "PACKING LIST NO" in upper:
            return "PACKING_LIST"
        if "CERTIFICATE OF ORIGIN" in upper or "FORM D" in upper or "FORM E" in upper:
            return "COO"
        if (
            "BILL OF LADING INSTRUCTION" in upper
            or "BL INSTRUCTION" in upper
            or "SHIPPING INSTRUCTION" in upper
            or "SHIPPING INSTRUCTIONS" in upper
            or "SHIPPER'S INSTRUCTION" in upper
        ):
            return "SI"
        if "BILL OF LADING" in upper or "DRAFT BL" in upper or "BILL OF LADING (DRAFT)" in upper:
            return "BL"

        # Filename fallback
        upper_fn = filename.upper()
        if "_SI" in upper_fn or "SI_" in upper_fn or "SHIPPING_INSTRUCTION" in upper_fn:
            return "SI"
        if "_BL" in upper_fn or "BL_" in upper_fn or "BILL_OF_LADING" in upper_fn:
            return "BL"
        return "UNKNOWN"
